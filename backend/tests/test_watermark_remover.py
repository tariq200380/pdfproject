"""Unit tests for WatermarkRemoverService across PDF, Word, Excel, PowerPoint, and Images."""

import pytest
from pathlib import Path
import pymupdf
import docx
import openpyxl
from pptx import Presentation
from PIL import Image

from backend.app.services.watermark_remover import watermark_remover


def test_pdf_watermark_removal(tmp_path):
    # 1. Create a sample PDF with "CONFIDENTIAL" text
    pdf_in = tmp_path / "sample_watermark.pdf"
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((50, 100), "This is legitimate contract content.")
    page.insert_text((150, 400), "CONFIDENTIAL", fontsize=36, color=(0.8, 0.2, 0.2))
    
    # Add watermark annotation
    annot = page.add_rect_annot(pymupdf.Rect(100, 350, 500, 450))
    annot.set_info(content="Watermark Stamp")
    annot.update()
    doc.save(str(pdf_in))
    doc.close()

    # 2. Run watermark removal
    pdf_out = tmp_path / "cleaned_document.pdf"
    result = watermark_remover.remove_pdf_watermark(pdf_in, pdf_out, watermark_text="CONFIDENTIAL")
    assert result.exists()

    # 3. Verify cleaned PDF
    cleaned_doc = pymupdf.open(str(pdf_out))
    cleaned_text = cleaned_doc[0].get_text()
    assert "This is legitimate contract content." in cleaned_text
    assert "CONFIDENTIAL" not in cleaned_text
    cleaned_doc.close()


def test_docx_watermark_removal(tmp_path):
    # 1. Create a Word document with header watermark
    docx_in = tmp_path / "sample_watermark.docx"
    doc = docx.Document()
    doc.add_paragraph("First paragraph of contract.")
    
    # Add watermark-like text in header
    header = doc.sections[0].header
    p_header = header.paragraphs[0]
    p_header.text = "CONFIDENTIAL WATERMARK"
    
    doc.save(str(docx_in))

    # 2. Run watermark removal
    docx_out = tmp_path / "cleaned_document.docx"
    result = watermark_remover.remove_docx_watermark(docx_in, docx_out, watermark_text="CONFIDENTIAL")
    assert result.exists()

    # 3. Verify cleaned docx
    cleaned_doc = docx.Document(str(docx_out))
    assert len(cleaned_doc.paragraphs) > 0
    assert "First paragraph of contract." in cleaned_doc.paragraphs[0].text


def test_xlsx_watermark_removal(tmp_path):
    # 1. Create sample Excel file
    xlsx_in = tmp_path / "sample_watermark.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financials"
    ws["A1"] = "Revenue"
    ws["B1"] = 50000
    wb.save(str(xlsx_in))

    # 2. Run watermark removal
    xlsx_out = tmp_path / "cleaned_document.xlsx"
    result = watermark_remover.remove_xlsx_watermark(xlsx_in, xlsx_out)
    assert result.exists()

    # 3. Verify Excel is intact
    cleaned_wb = openpyxl.load_workbook(str(xlsx_out))
    cleaned_ws = cleaned_wb["Financials"]
    assert cleaned_ws["A1"].value == "Revenue"
    assert cleaned_ws["B1"].value == 50000
    cleaned_wb.close()


def test_pptx_watermark_removal(tmp_path):
    # 1. Create sample PPTX file
    pptx_in = tmp_path / "sample_watermark.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    title.text = "Presentation Title"
    
    # Add a watermark shape
    txBox = slide.shapes.add_textbox(100, 100, 300, 100)
    tf = txBox.text_frame
    tf.text = "DRAFT WATERMARK"
    
    prs.save(str(pptx_in))

    # 2. Run watermark removal
    pptx_out = tmp_path / "cleaned_presentation.pptx"
    result = watermark_remover.remove_pptx_watermark(pptx_in, pptx_out, watermark_text="DRAFT")
    assert result.exists()

    # 3. Verify PPTX
    cleaned_prs = Presentation(str(pptx_out))
    slide = cleaned_prs.slides[0]
    texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
    assert "Presentation Title" in texts
    assert not any("DRAFT" in t for t in texts)


def test_image_watermark_removal(tmp_path):
    # 1. Create a sample image
    img_in = tmp_path / "sample_watermark.png"
    img = Image.new("RGB", (200, 200), color=(255, 255, 255))
    img.save(str(img_in))

    # 2. Run watermark removal
    img_out = tmp_path / "cleaned_image.png"
    result = watermark_remover.remove_image_watermark(img_in, img_out)
    assert result.exists()


@pytest.mark.asyncio
async def test_api_remove_watermark_endpoint(tmp_path):
    from httpx import ASGITransport, AsyncClient
    from backend.app.main import app

    # Create dummy PDF with watermark
    doc = pymupdf.open()
    page = doc.new_page(width=400, height=400)
    page.insert_text((50, 50), "Real Document Text")
    page.insert_text((100, 200), "CONFIDENTIAL", fontsize=30)
    pdf_bytes = doc.tobytes()
    doc.close()

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        files = {"file": ("test_wm.pdf", pdf_bytes, "application/pdf")}
        data = {"watermark_text": "CONFIDENTIAL"}
        res = await ac.post("/api/pdf/remove-watermark", files=files, data=data)
        assert res.status_code == 200
        assert res.headers["content-type"] == "application/pdf"
        assert len(res.content) > 0
        
        # Verify content has removed watermark
        cleaned_doc = pymupdf.open("pdf", res.content)
        text = cleaned_doc[0].get_text()
        assert "Real Document Text" in text
        assert "CONFIDENTIAL" not in text
        cleaned_doc.close()


def test_video_watermark_removal_9_16_and_16_9(tmp_path):
    import subprocess

    # 1. Generate 9:16 Vertical test video (TikTok / Reels style: 720x1280)
    video_9_16 = tmp_path / "test_reels_9_16.mp4"
    subprocess.run([
        "ffmpeg", "-y", "-f", "lavfi", "-i", "color=c=black:s=720x1280:d=1",
        "-c:v", "libx264", "-crf", "20", str(video_9_16)
    ], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

    out_9_16 = tmp_path / "cleaned_reels_9_16.mp4"
    res_9_16 = watermark_remover.remove_video_watermark(
        input_path=video_9_16,
        output_path=out_9_16,
        position="tiktok-dual",
        aspect_ratio="9:16",
    )
    assert res_9_16.exists()
    assert res_9_16.stat().st_size > 0

    # 2. Generate 16:9 Landscape test video (YouTube style: 1280x720)
    video_16_9 = tmp_path / "test_youtube_16_9.mp4"
    subprocess.run([
        "ffmpeg", "-y", "-f", "lavfi", "-i", "color=c=black:s=1280x720:d=1",
        "-c:v", "libx264", "-crf", "20", str(video_16_9)
    ], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

    out_16_9 = tmp_path / "cleaned_youtube_16_9.mp4"
    res_16_9 = watermark_remover.remove_video_watermark(
        input_path=video_16_9,
        output_path=out_16_9,
        position="bottom-right",
        aspect_ratio="16:9",
    )
    assert res_16_9.exists()
    assert res_16_9.stat().st_size > 0


